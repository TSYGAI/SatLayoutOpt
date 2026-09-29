"""文件说明：从布局位姿、历史质量和本体几何生成安装参数 Excel；不使用安全包络尺寸，不修改任务或模型库。

本体几何优先读取任务快照；旧任务只允许从摘要一致的固化 STEP 恢复，依赖现有 STEP 离散器和 openpyxl。
"""

import hashlib
from io import BytesIO
import math

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.core.config import runtime_config
from app.domain.layout_config import LAYOUT_CONFIG
from app.domain.layout_geometry import rotation_z
from app.domain.model_envelope import load_step_triangle_mesh


def load_body_vertices(directory, placement, model_store):
    """返回源模型坐标下的本体顶点；旧快照回读并校验原始 STEP，缺失或摘要不同抛 ValueError。"""
    snapshot = directory / f"{placement['record_id']}.npz"
    if snapshot.is_file():
        with np.load(snapshot, allow_pickle=False) as data:
            if "body_vertices" in data:
                return np.array(data["body_vertices"], dtype=float)
    record = model_store.get_record(placement["record_id"])
    if record is None or record["source"]["sha256"] != placement["source_sha256"]:
        raise ValueError(f"{placement['name']} 的历史任务未保存本体快照，且原始模型已缺失或改变，无法导出本体尺寸")
    try:
        source = model_store.get_asset_path(record["record_id"], record["source"]["asset_name"])
        with source.open("rb") as stream:
            digest = hashlib.file_digest(stream, "sha256").hexdigest()
        if digest != placement["source_sha256"]:
            raise ValueError(f"{placement['name']} 的原始模型摘要不一致，无法导出本体尺寸")
        mesh = load_step_triangle_mesh(
            source,
            linear_ratio=runtime_config.model_envelope_step_linear_ratio,
            minimum_linear_tolerance_mm=runtime_config.model_envelope_step_min_linear_tolerance_mm,
            angular_degrees=runtime_config.model_envelope_step_angular_degrees,
        )
    except (FileNotFoundError, LookupError) as error:
        raise ValueError(f"{placement['name']} 的原始模型文件已缺失，无法导出本体尺寸") from error
    return np.asarray(mesh.vertices, dtype=float)


def installation_row(placement, body_vertices):
    """返回本体尺寸、安装角度、质心高度、五项默认值及单星重量；单位 mm/度/kg，无效几何或质量抛 ValueError。

    长宽高按安装姿态去掉环向旋转后的 X/Y/Z 外接尺寸计算，保留实际单位、缩放和安装面调整。
    侧挂方位取安装基准点的 atan2(y, x)；居中顶部无环向安装位置，约定为零度。
    重量只读取方案保存的 mass_kg，不从模型库当前设置回填历史质量。
    """
    try:
        mass = float(placement["mass_kg"])
    except (KeyError, TypeError, ValueError) as error:
        raise ValueError(f"{placement['name']} 的历史重量缺失或无效，无法导出") from error
    if not math.isfinite(mass) or mass <= 0 or isinstance(placement["mass_kg"], bool):
        raise ValueError(f"{placement['name']} 的历史重量缺失或无效，无法导出")
    vertices = np.asarray(body_vertices, dtype=float)
    matrix = np.asarray(placement["model_to_global_matrix"], dtype=float)
    center = np.asarray(placement["center_of_mass_mm"], dtype=float)
    face_center = np.asarray(placement["installation_face_center_mm"], dtype=float)
    angle = float(placement["rotation_z_deg"])
    if (vertices.ndim != 2 or vertices.shape[1] != 3 or not len(vertices)
            or matrix.shape != (4, 4) or center.shape != (3,) or face_center.shape != (3,)
            or not all(np.isfinite(value).all() for value in (vertices, matrix, center, face_center))
            or not math.isfinite(angle)):
        raise ValueError(f"{placement['name']} 的本体或安装位姿无效，无法导出")
    local_rotation = rotation_z(math.radians(-angle)) @ matrix[:3, :3]
    dimensions = np.ptp(vertices @ local_rotation.T, axis=0)
    if np.any(dimensions <= 0):
        raise ValueError(f"{placement['name']} 的本体尺寸无效，无法导出")
    if placement["mount"] == "top":
        bearing = 0.0
    else:
        if np.linalg.norm(face_center[:2]) <= LAYOUT_CONFIG["geometry_tolerance_mm"]:
            raise ValueError(f"{placement['name']} 的侧挂安装方位无法确定")
        bearing = math.degrees(math.atan2(face_center[1], face_center[0])) % 360.0
    defaults = LAYOUT_CONFIG["export"]
    return [*map(float, dimensions), bearing, float(center[2]),
            defaults["length_mm"], defaults["height_mm"], defaults["width_mm"],
            defaults["horizontal_mm"], defaults["vertical_mm"], mass]


def export_candidate_workbook(job, candidate_id, directory, model_store):
    """导出指定有效候选为内存 XLSX 字节；保持 placement 顺序，不改变已选方案，未知候选抛 LookupError。"""
    candidate = next((item for item in job["candidates"] if item["candidate_id"] == candidate_id), None)
    if candidate is None:
        raise LookupError("找不到指定的有效布局方案")
    if not candidate.get("validation", {}).get("valid") or not candidate["placements"]:
        raise ValueError("当前方案尚未通过布局验收，无法导出")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "安装参数"
    sheet.append(["", "长", "宽", "高", "安装角度", "安装高度", "长度", "高度", "宽度", "水平", "竖直", "重量（kg）"])
    bodies = {}
    for index, placement in enumerate(candidate["placements"], start=1):
        key = (placement["record_id"], placement["source_sha256"])
        if key not in bodies:
            bodies[key] = load_body_vertices(directory, placement, model_store)
        sheet.append([index, *installation_row(placement, bodies[key])])
    sheet.freeze_panes = "B2"
    for row in sheet:
        for cell in row:
            cell.font = Font(name="宋体", size=12, bold=cell.row == 1)
            cell.alignment = Alignment(horizontal="left" if cell.row == 1 else "right", vertical="center")
            if cell.row > 1:
                cell.number_format = "0.000" if 2 <= cell.column <= 5 else "0.###" if cell.column in (6, sheet.max_column) else "0"
        sheet.row_dimensions[row[0].row].height = 23
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 7 if column == 1 else 16
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()
